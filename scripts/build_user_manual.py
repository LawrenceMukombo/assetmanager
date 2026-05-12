#!/usr/bin/env python3
"""Generate the NPAMS User Manual in both .docx and .pdf formats."""
from __future__ import annotations
import os
from datetime import date

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, PageBreak,
    Table, TableStyle, ListFlowable, ListItem, KeepTogether,
)
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_JUSTIFY

ICSA_BLUE = "#0F4C81"
ICSA_RED = "#CE1126"
ICSA_GOLD = "#FCD116"
TODAY = date.today().strftime("%d %B %Y")

# ---------------------------------------------------------------------------
# Manual content — single source of truth for both formats.
# A section is (heading, [blocks]) where each block is one of:
#   ("p",  "paragraph text")
#   ("ul", ["bullet 1", "bullet 2", ...])
#   ("ol", ["step 1", "step 2", ...])
#   ("kv", [("Field", "Description"), ...])      -> two-column table
#   ("tbl", [["h1","h2","h3"], ["r1c1",...], ...]) -> table with header row
#   ("note", "callout text")
# ---------------------------------------------------------------------------

ROLES_TABLE = [
    ["Role", "Scope", "Typical Responsibilities"],
    ["Super Admin", "Whole system", "Configure tenants, agencies, roles; full access to every module."],
    ["National Asset Controller", "National", "Approve high-value purchases, run national reports, oversee audits across provinces."],
    ["Provincial Admin", "Single province", "Manage users, assets, stock and purchase requests within their province."],
    ["Agency Admin", "Single agency (e.g. ICSA)", "Manage agency-specific assets, stock catalogue, purchases and maintenance."],
    ["Provincial Asset Officer", "Province / Facility", "Field officer — record verifications, raise purchase requests, log maintenance."],
    ["National Auditor", "Read-only national", "Run audit sessions, verify items, sign off audit reports."],
    ["Viewer", "As scoped", "Read-only access for stakeholders and observers."],
]

PR_STATUSES_TABLE = [
    ["Status", "Meaning"],
    ["Draft", "Saved but not yet submitted; only the requester can edit or submit."],
    ["Submitted", "Awaiting approval by an Admin or National Asset Controller."],
    ["Approved", "Cleared for procurement; quantities can now be received against it."],
    ["Rejected", "Returned with a reason; no further action possible (raise a new PR)."],
    ["Received", "Goods fully received; stock balances were updated automatically."],
    ["Closed", "Archived; no further changes."],
]

MAINTENANCE_PRIORITY_TABLE = [
    ["Priority", "When to use"],
    ["Critical", "Asset is offline or unsafe; must be addressed within 48 hours."],
    ["High", "Service overdue or partial failure; address within 7 days."],
    ["Medium", "Routine scheduled service or non-blocking issue; address within 30 days."],
    ["Low", "Minor or cosmetic issue; can be batched with other work."],
]

AUDIT_ITEM_STATUS_TABLE = [
    ["Item Status", "Meaning"],
    ["Pending", "Not yet sighted by the auditor."],
    ["Verified", "Sighted in person; condition confirmed and (optionally) GPS / photo captured."],
    ["Not Found", "Asset is missing from its registered location at audit time."],
    ["Damaged", "Asset is present but damaged beyond normal wear; flagged for follow-up."],
]

ASSET_STATUS_TABLE = [
    ["Status", "Meaning"],
    ["Active", "In service at its registered location."],
    ["Under Maintenance", "Temporarily out of service for repair or scheduled maintenance."],
    ["Missing", "Cannot be located; under investigation."],
    ["Disposed", "Permanently removed from the register (sold, scrapped, written off)."],
]

SECTIONS = [
    ("1. Introduction", [
        ("p",
         "The National Public Asset Management System (NPAMS) is the official platform "
         "for tracking the assets, stock and procurement activity of Papua New Guinea "
         "public service agencies. This manual covers the day-to-day workflows for "
         "officers, administrators and auditors using the system."),
        ("p",
         "NPAMS is currently deployed for the PNG Immigration & Citizenship Service "
         "Authority (ICSA) and is designed to scale to additional agencies. Where this "
         "manual mentions ICSA-specific data, the same workflows apply to any agency "
         "configured in the platform."),
        ("p", "This manual was prepared on " + TODAY + ". Software by LanFrame."),
    ]),

    ("2. Getting Started", [
        ("p",
         "NPAMS is a web application. Open the URL provided by your administrator in "
         "any modern browser (Chrome, Edge, Safari, Firefox). On a phone or tablet the "
         "interface adapts to a single-column layout."),
        ("p", "Default seeded administrator credentials (change on first login):"),
        ("kv", [
            ("Email", "immigration.admin@npams.gov.pg"),
            ("Password", "Admin1234!"),
        ]),
        ("note",
         "Change the default password immediately after the first sign-in. Use Settings "
         "→ Account to set a new password and update your contact details."),
        ("p", "If you forget your password, click \"Forgot password?\" on the login screen. "
         "An email with a one-time reset link will be sent to your registered address. "
         "Reset links expire after one hour and can only be used once."),
    ]),

    ("3. Roles and Permissions", [
        ("p",
         "NPAMS uses role-based access control with geographic and agency scoping. "
         "What you can see and change is the intersection of (a) your role and "
         "(b) your scope (national, provincial, agency or facility). Officers cannot "
         "see records outside their scope, and administrators can only manage users "
         "within their scope."),
        ("tbl", ROLES_TABLE),
        ("p",
         "If a button or menu item is not visible to you, it is because your role does "
         "not have permission for that action. Speak to your administrator if you "
         "believe you need additional access."),
    ]),

    ("4. The Dashboard", [
        ("p",
         "The Dashboard is your landing page after sign-in. It summarises the assets, "
         "stock and workflow items that fall within your scope and highlights items "
         "that need attention."),
        ("ul", [
            "Total assets, with breakdowns by status and condition.",
            "Total inventory value (book value, net of depreciation).",
            "Low-stock alerts — items below reorder level at one or more locations.",
            "Pending purchase requests awaiting your approval.",
            "Upcoming and overdue maintenance.",
            "Active audit sessions and your assignments within them.",
            "Recent activity feed across the modules you can see.",
        ]),
    ]),

    ("5. Assets", [
        ("p",
         "The Assets module is the core register. Every physical asset that the "
         "agency owns — buildings, vehicles, ICT equipment, biometric stations, "
         "furniture, uniforms — is recorded here with a unique asset tag."),

        ("p", "Browsing assets:"),
        ("ol", [
            "Open Assets from the left navigation.",
            "Use the search box to find an asset by tag, name, serial number or supplier.",
            "Use the filter chips for category, status, condition, province or facility.",
            "Click any row to open the asset detail page.",
            "Use the page-size selector at the bottom to show 10, 25, 50 or 100 rows per page.",
        ]),

        ("p", "Creating an asset (Admin or Officer roles):"),
        ("ol", [
            "Click + New Asset on the Assets page.",
            "Fill in the asset tag (must be unique), name and category.",
            "Record purchase information: date, cost, supplier and useful life in years.",
            "Set the location (province, district, facility) and assign a custodian if known.",
            "Set Depreciation Method (straight-line by default) and salvage value.",
            "Optionally upload a photo and add notes.",
            "Click Save. The asset is now visible to anyone within its scope.",
        ]),

        ("p", "Asset statuses:"),
        ("tbl", ASSET_STATUS_TABLE),

        ("p", "Asset transfers:"),
        ("p",
         "From an asset detail page, click Transfer to move an asset between facilities "
         "or custodians. Transfers are logged in the asset history with the date, "
         "from/to locations and the user who performed the transfer."),

        ("p", "Public asset lookup:"),
        ("p",
         "Each asset has a public verification page accessible via its asset tag. "
         "This page shows only non-sensitive information (tag, category, custodian "
         "agency, last verified date) and is suitable for QR-code labels in the field."),
    ]),

    ("6. Categories", [
        ("p",
         "Categories group similar assets and drive the depreciation defaults and the "
         "Stock dashboard. Examples used by ICSA include Buildings & Infrastructure, "
         "Vehicles & Transport, ICT Equipment, Border Control Equipment, Biometric & "
         "Identity Capture, Passport & Document Production, Communication Equipment, "
         "Office Furniture, and Uniforms & Accoutrements."),
        ("p",
         "Only Super Admins and National Asset Controllers can add, rename or "
         "deactivate categories. Renaming a category updates every asset that uses it "
         "in real time."),
    ]),

    ("7. Stock Management", [
        ("p",
         "Stock covers consumable inventory — paper, toner, vehicle spares, uniform "
         "issues, stationery — that is held at one or more facilities and topped up "
         "from purchase requests."),

        ("p", "Stock items vs stock balances:"),
        ("ul", [
            "A stock item is the SKU itself (item code, name, unit of measure, default reorder level).",
            "A stock balance is the quantity of that item held at a specific facility.",
            "An item with no per-location balance falls back to its agency-wide on-hand quantity.",
        ]),

        ("p", "Low-stock alerts:"),
        ("p",
         "An item is flagged \"low\" when any of its per-location balances sits at or "
         "below its reorder threshold. The Stock dashboard highlights these in red and "
         "the home Dashboard surfaces them in the alerts panel."),

        ("p", "Recording stock movements:"),
        ("ol", [
            "Open the stock item from the Stock list.",
            "Click Adjust Stock and choose Issue, Receive, Transfer or Adjust.",
            "Select the source / destination facility and quantity.",
            "Add a note describing why (e.g. \"Q1 issue to Vanimo border post\").",
            "Save. The relevant balances update automatically and the movement appears in the audit log.",
        ]),
    ]),

    ("8. Purchase Requests", [
        ("p",
         "Purchase Requests (PRs) capture the full procurement workflow — from a "
         "request being raised by a custodian, to approval by management, to receipt "
         "of goods and update of stock balances."),

        ("p", "PR lifecycle:"),
        ("tbl", PR_STATUSES_TABLE),

        ("p", "Raising a PR:"),
        ("ol", [
            "Open Purchase Requests and click + New Request.",
            "Pick the stock item and quantity required.",
            "Enter the supplier, unit cost (estimate) and any notes (justification, urgency).",
            "Save as Draft (to come back later) or click Submit to send for approval.",
        ]),

        ("p", "Approving a PR:"),
        ("ol", [
            "Open the Purchase Requests list and filter by status = Submitted.",
            "Click the PR to open its detail page.",
            "Review the request, supplier and cost.",
            "Click Approve or Reject. If rejecting, you must provide a reason — this is recorded permanently and is visible to the requester.",
        ]),

        ("p", "Receiving goods against a PR:"),
        ("ol", [
            "Open the approved PR.",
            "Click Receive Goods and enter the quantity received (may be partial).",
            "Confirm. The receiving facility's stock balance is incremented automatically and an event is added to the PR's signed audit trail.",
            "Once the full quantity is received, the PR moves to Received. Closing it is optional — Closed simply archives it.",
        ]),

        ("note",
         "Every state change on a PR (submit, approve, reject, receive, close) is "
         "signed and timestamped with the acting user. This trail cannot be edited "
         "or deleted and is shown on the PR detail page."),
    ]),

    ("9. Maintenance", [
        ("p",
         "The Maintenance module schedules and tracks preventative and corrective "
         "maintenance for assets — vehicle services, printer repairs, eGate firmware "
         "upgrades, generator overhauls, and so on."),

        ("p", "Priorities:"),
        ("tbl", MAINTENANCE_PRIORITY_TABLE),

        ("p", "Scheduling maintenance:"),
        ("ol", [
            "Open Maintenance and click + Schedule Maintenance.",
            "Choose the asset (search by tag or name).",
            "Enter a clear title (e.g. \"Engine service — overdue\") and description.",
            "Set the priority and the scheduled date.",
            "Optionally assign a responsible officer and an estimated cost.",
            "Save. The asset's status will show Under Maintenance once the work begins.",
        ]),

        ("p", "Completing maintenance:"),
        ("ol", [
            "Open the maintenance record and click Mark Complete.",
            "Enter the actual cost and completion notes.",
            "Save. The asset returns to Active status and the cost is reflected in the asset's lifetime maintenance total.",
        ]),
    ]),

    ("10. Audit Sessions", [
        ("p",
         "Audits are the periodic verification of physical assets against the register. "
         "An audit session covers a date range, a province (or national) and a set of "
         "facilities. Within each session, assignments are issued to auditors and each "
         "asset within an assignment is verified individually."),

        ("p", "Audit hierarchy:"),
        ("ul", [
            "Session — the umbrella for the audit (e.g. \"Q2 2026 National Asset Verification\").",
            "Assignment — one auditor's responsibility for one facility within the session.",
            "Item — a single asset within an assignment, verified by the auditor.",
        ]),

        ("p", "Audit item statuses:"),
        ("tbl", AUDIT_ITEM_STATUS_TABLE),

        ("p", "Verifying an asset in the field:"),
        ("ol", [
            "Open Audit, then your assignment.",
            "Tap an asset that is set to Pending.",
            "Use the Verify screen to capture: condition observed, GPS coordinates (auto from device), and a photo of the asset.",
            "Add notes if anything differs from the register.",
            "Tap Submit. The item is signed by you and timestamped.",
            "Continue through the assignment. The session can be closed once all items have been verified or marked Not Found / Damaged.",
        ]),

        ("note",
         "GPS and photo capture require the browser permission to be granted on first "
         "use. If the device denies it, the item can still be verified manually with a "
         "note explaining the absence of GPS or photo evidence."),
    ]),

    ("11. Locations", [
        ("p",
         "Locations are the geographic scaffolding of NPAMS: provinces contain "
         "districts, which contain facilities. Every asset, user and stock balance is "
         "linked to one of these levels for scoping."),
        ("ul", [
            "Provinces are the 22 PNG provinces and are managed centrally.",
            "Districts are the LLG-equivalent areas within each province.",
            "Facilities are the individual ICSA offices, border posts, sea ports and airport posts.",
        ]),
        ("p",
         "The Locations page lets administrators add, rename or deactivate facilities. "
         "Deactivating a facility hides it from selection in new records but preserves "
         "history for any asset that was previously located there."),
        ("p",
         "The GIS map view (left navigation → Map) plots every facility on a map of "
         "Papua New Guinea, with colour-coded markers showing asset counts and any "
         "open audit assignments."),
    ]),

    ("12. Users and Access", [
        ("p",
         "User Administration is available to Super Admins, National Asset Controllers, "
         "Provincial Admins and Agency Admins. Each can only create and edit users "
         "within their own scope."),

        ("p", "Creating a user:"),
        ("ol", [
            "Open Users → + New User.",
            "Enter full name, work email, role and home location (province / district / facility, or agency).",
            "Set an initial password — the user will be prompted to change it on first sign-in.",
            "Save. The user receives an email with their sign-in details.",
        ]),

        ("p", "Resetting a user's password:"),
        ("ol", [
            "Open the user's profile in the Users list.",
            "Click Send Password Reset Link. The user will receive a one-time email link.",
            "Alternatively, click Set Temporary Password to choose a temporary password yourself.",
        ]),
        ("note",
         "A repeated burst of failed password resets for the same account triggers an "
         "alert badge on the user record. Investigate before issuing further resets."),

        ("p", "Deactivating a user:"),
        ("p",
         "From the user's profile, toggle Active off. The user can no longer sign in, "
         "but their historical actions (PR approvals, audit verifications, etc.) are "
         "preserved and remain attributable to them."),
    ]),

    ("13. Reports", [
        ("p",
         "Reports brings together printable, exportable summaries of the data in NPAMS:"),
        ("ul", [
            "Asset register — full list with filters; export to CSV or PDF.",
            "Depreciation schedule — current book value vs original cost per asset and per category.",
            "Stock-on-hand — quantities by item and facility, with low-stock highlighting.",
            "Purchase request summary — counts and value by status, supplier and requester.",
            "Maintenance log — completed and outstanding work, with cost roll-up.",
            "Audit completion — verification rates and exception lists per session.",
        ]),
        ("p",
         "All exports respect your scope — you will only see the rows you are allowed "
         "to see in the rest of the system."),
    ]),

    ("14. Notifications", [
        ("p",
         "The bell icon in the top-right surfaces in-app notifications: new purchase "
         "requests awaiting your approval, maintenance becoming overdue, audit "
         "assignments issued to you, and password reset alerts on accounts you "
         "administer. Click a notification to jump to the related record."),
    ]),

    ("15. Settings and System Status", [
        ("p", "Settings (top-right user menu) covers your personal account:"),
        ("ul", [
            "Update your name, contact phone and home location.",
            "Change your password.",
            "Configure email notification preferences.",
        ]),
        ("p",
         "System Status (Super Admin only) shows the live health of the API server, "
         "background workers, database and email subsystem, plus the last successful "
         "auto-seed run. Use this page first when investigating any platform issue."),
    ]),

    ("16. Data Retention and Audit Trail", [
        ("p",
         "NPAMS preserves a tamper-evident history for every state change on the core "
         "modules. In particular:"),
        ("ul", [
            "Purchase request events (submit, approve, reject, receive, close) are signed and timestamped per user.",
            "Asset transfers, depreciation method changes and disposals are logged on the asset history tab.",
            "Audit item verifications include the signing user, timestamp, GPS and photo where provided.",
            "User activity (logins, password resets, role changes) appears in the global activity log accessible to administrators.",
        ]),
        ("p",
         "These records are append-only — they cannot be edited or deleted from the "
         "user interface."),
    ]),

    ("17. Frequently Asked Questions", [
        ("kv", [
            ("I cannot see a facility / asset that I expect to see.",
             "Your role is scoped to a province, agency or facility. Ask your administrator to confirm your scope."),
            ("My password reset email has not arrived.",
             "Check your spam folder. Reset links also expire after one hour — request a new one if necessary."),
            ("I approved a PR by mistake.",
             "PR state changes cannot be reversed. Add a closing note and create a new PR if the goods should not be procured."),
            ("Why does an asset still show as Under Maintenance after I marked the work complete?",
             "Refresh the page. The status updates in real time once the maintenance record is saved as Complete."),
            ("Can I export everything to Excel?",
             "Yes — the Reports module produces CSV exports that open directly in Excel or Google Sheets."),
            ("How long does my session stay signed in?",
             "Sessions remain valid for 24 hours of activity. After that you will be asked to sign in again."),
        ]),
    ]),

    ("18. Glossary", [
        ("kv", [
            ("Agency", "An employing organisation in NPAMS (e.g. ICSA). Each agency owns its own assets and stock."),
            ("Asset Tag", "The unique human-readable identifier printed on the asset label (e.g. PNGICA-VEH-001)."),
            ("Audit Session", "A scheduled physical verification campaign covering a set of facilities and a date range."),
            ("Custodian", "The user accountable for an asset's day-to-day care."),
            ("Depreciation", "The reduction in an asset's book value over its useful life. NPAMS uses straight-line by default."),
            ("Facility", "A specific physical location — office, border post, sea port, airport post, etc."),
            ("ICSA", "The PNG Immigration & Citizenship Service Authority — first agency live on NPAMS."),
            ("Purchase Request (PR)", "A formal request to procure stock; subject to an approval and receipt workflow."),
            ("Reorder Level", "The threshold at which a stock balance is flagged as low and a PR should be raised."),
            ("Salvage Value", "The estimated residual value of an asset at the end of its useful life."),
            ("Scope", "The geographic / organisational subset of records a user is allowed to see and change."),
            ("Stock Balance", "The quantity of a stock item held at a specific facility."),
            ("Useful Life", "The number of years over which an asset is depreciated."),
        ]),
    ]),

    ("19. Support", [
        ("p",
         "For account access and routine \"how do I…\" questions, contact your "
         "agency administrator in the first instance. For platform faults — pages "
         "that fail to load, data that appears corrupted, or suspected security "
         "incidents — escalate to the National Asset Controller, who can engage "
         "the LanFrame support team."),
        ("p",
         "When raising a support ticket, include: your full name, role, the URL of "
         "the page where the issue occurred, the exact time, and a screenshot if "
         "possible. This information dramatically speeds up diagnosis."),
    ]),
]


# ---------------------------------------------------------------------------
# DOCX builder
# ---------------------------------------------------------------------------

def _set_cell_shade(cell, hex_color: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color.lstrip("#"))
    tcPr.append(shd)


def _hex_rgb(hex_color: str) -> RGBColor:
    h = hex_color.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def build_docx(out_path: str) -> None:
    doc = Document()

    # Set default body font
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    section = doc.sections[0]
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)

    # ── Cover ──
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("NPAMS")
    run.bold = True
    run.font.size = Pt(48)
    run.font.color.rgb = _hex_rgb(ICSA_BLUE)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("National Public Asset Management System")
    run.font.size = Pt(18)
    run.font.color.rgb = _hex_rgb(ICSA_BLUE)

    doc.add_paragraph()
    band = doc.add_paragraph()
    band.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = band.add_run("USER MANUAL")
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = _hex_rgb(ICSA_RED)

    for _ in range(8):
        doc.add_paragraph()

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run(
        "Prepared for the PNG Immigration & Citizenship Service Authority\n"
        f"Issued {TODAY}\nSoftware by LanFrame"
    )
    run.font.size = Pt(12)
    run.font.color.rgb = _hex_rgb(ICSA_BLUE)

    doc.add_page_break()

    # ── Table of contents ──
    h = doc.add_paragraph()
    run = h.add_run("Contents")
    run.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = _hex_rgb(ICSA_BLUE)
    for heading, _ in SECTIONS:
        p = doc.add_paragraph(heading)
        p.paragraph_format.space_after = Pt(2)
    doc.add_page_break()

    # ── Body ──
    for heading, blocks in SECTIONS:
        h = doc.add_paragraph()
        run = h.add_run(heading)
        run.bold = True
        run.font.size = Pt(18)
        run.font.color.rgb = _hex_rgb(ICSA_BLUE)
        h.paragraph_format.space_before = Pt(6)
        h.paragraph_format.space_after = Pt(6)

        for block in blocks:
            kind = block[0]
            payload = block[1]

            if kind == "p":
                p = doc.add_paragraph(payload)
                p.paragraph_format.space_after = Pt(8)

            elif kind == "ul":
                for item in payload:
                    doc.add_paragraph(item, style="List Bullet")

            elif kind == "ol":
                for item in payload:
                    doc.add_paragraph(item, style="List Number")

            elif kind == "kv":
                t = doc.add_table(rows=len(payload), cols=2)
                t.autofit = False
                for i, (k, v) in enumerate(payload):
                    cell_k, cell_v = t.rows[i].cells
                    cell_k.width = Inches(2.2)
                    cell_v.width = Inches(4.3)
                    pk = cell_k.paragraphs[0]
                    rk = pk.add_run(k)
                    rk.bold = True
                    rk.font.color.rgb = _hex_rgb(ICSA_BLUE)
                    cell_v.paragraphs[0].add_run(v)
                doc.add_paragraph()

            elif kind == "tbl":
                rows = payload
                t = doc.add_table(rows=len(rows), cols=len(rows[0]))
                for j, header in enumerate(rows[0]):
                    cell = t.rows[0].cells[j]
                    _set_cell_shade(cell, ICSA_BLUE)
                    cell.paragraphs[0].text = ""
                    r = cell.paragraphs[0].add_run(header)
                    r.bold = True
                    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                for i, row in enumerate(rows[1:], start=1):
                    for j, val in enumerate(row):
                        t.rows[i].cells[j].text = val
                doc.add_paragraph()

            elif kind == "note":
                p = doc.add_paragraph()
                run = p.add_run("Note: " + payload)
                run.italic = True
                run.font.color.rgb = _hex_rgb(ICSA_RED)
                p.paragraph_format.space_after = Pt(10)

        doc.add_paragraph()

    doc.save(out_path)


# ---------------------------------------------------------------------------
# PDF builder
# ---------------------------------------------------------------------------

def build_pdf(out_path: str) -> None:
    blue = colors.HexColor(ICSA_BLUE)
    red = colors.HexColor(ICSA_RED)
    gold = colors.HexColor(ICSA_GOLD)

    base = getSampleStyleSheet()
    body = ParagraphStyle(
        "body", parent=base["BodyText"], fontName="Helvetica", fontSize=10.5,
        leading=15, spaceAfter=6, alignment=TA_JUSTIFY,
    )
    h1 = ParagraphStyle(
        "h1", parent=base["Heading1"], fontName="Helvetica-Bold", fontSize=16,
        leading=20, textColor=blue, spaceBefore=10, spaceAfter=8,
    )
    title_style = ParagraphStyle(
        "title", parent=base["Title"], fontName="Helvetica-Bold", fontSize=44,
        leading=52, textColor=blue, alignment=TA_CENTER,
    )
    subtitle = ParagraphStyle(
        "subtitle", parent=base["Title"], fontName="Helvetica", fontSize=18,
        leading=22, textColor=blue, alignment=TA_CENTER, spaceAfter=10,
    )
    band = ParagraphStyle(
        "band", parent=base["Title"], fontName="Helvetica-Bold", fontSize=20,
        leading=24, textColor=red, alignment=TA_CENTER, spaceBefore=20,
    )
    meta = ParagraphStyle(
        "meta", parent=body, fontSize=11, alignment=TA_CENTER, textColor=blue,
    )
    note = ParagraphStyle(
        "note", parent=body, fontSize=10.5, leftIndent=12, rightIndent=12,
        textColor=red, fontName="Helvetica-Oblique", spaceBefore=4, spaceAfter=10,
    )
    toc_entry = ParagraphStyle(
        "toc", parent=body, fontSize=11, spaceAfter=2, alignment=TA_LEFT,
    )

    story: list = []

    # ── Cover ──
    story.append(Spacer(1, 4 * cm))
    story.append(Paragraph("NPAMS", title_style))
    story.append(Paragraph("National Public Asset Management System", subtitle))
    story.append(Spacer(1, 1 * cm))
    story.append(Paragraph("USER MANUAL", band))
    story.append(Spacer(1, 5 * cm))
    story.append(Paragraph(
        "Prepared for the PNG Immigration &amp; Citizenship Service Authority<br/>"
        f"Issued {TODAY}<br/>Software by LanFrame",
        meta,
    ))
    story.append(PageBreak())

    # ── TOC ──
    story.append(Paragraph("Contents", h1))
    for heading, _ in SECTIONS:
        story.append(Paragraph(heading, toc_entry))
    story.append(PageBreak())

    # ── Body ──
    for heading, blocks in SECTIONS:
        story.append(Paragraph(heading, h1))
        for block in blocks:
            kind = block[0]
            payload = block[1]

            if kind == "p":
                story.append(Paragraph(payload, body))

            elif kind == "ul":
                items = [ListItem(Paragraph(t, body), leftIndent=10) for t in payload]
                story.append(ListFlowable(items, bulletType="bullet", leftIndent=14))
                story.append(Spacer(1, 4))

            elif kind == "ol":
                items = [ListItem(Paragraph(t, body), leftIndent=10) for t in payload]
                story.append(ListFlowable(items, bulletType="1", leftIndent=14))
                story.append(Spacer(1, 4))

            elif kind == "kv":
                rows = [[Paragraph(f"<b>{k}</b>", body), Paragraph(v, body)] for k, v in payload]
                t = Table(rows, colWidths=[5.5 * cm, 11 * cm])
                t.setStyle(TableStyle([
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("BOX", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")),
                    ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#e2e8f0")),
                    ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#f1f5f9")),
                    ("LEFTPADDING", (0, 0), (-1, -1), 6),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                    ("TOPPADDING", (0, 0), (-1, -1), 5),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                ]))
                story.append(KeepTogether(t))
                story.append(Spacer(1, 8))

            elif kind == "tbl":
                rows = payload
                ncols = len(rows[0])
                # Allocate widths proportionally based on max col content length.
                weights = [max(len(str(r[i])) for r in rows) for i in range(ncols)]
                total_w = 16.5 * cm
                weight_sum = sum(weights) or 1
                col_widths = [total_w * (w / weight_sum) for w in weights]

                wrapped = [[Paragraph(str(c), body) for c in row] for row in rows]
                t = Table(wrapped, colWidths=col_widths, repeatRows=1)
                t.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), blue),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
                    ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#e2e8f0")),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
                    ("LEFTPADDING", (0, 0), (-1, -1), 6),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                    ("TOPPADDING", (0, 0), (-1, -1), 5),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                ]))
                story.append(t)
                story.append(Spacer(1, 8))

            elif kind == "note":
                story.append(Paragraph("<b>Note:</b> " + payload, note))

        story.append(Spacer(1, 6))

    def _on_page(canvas, doc):
        canvas.saveState()
        # Top accent band
        canvas.setFillColor(blue)
        canvas.rect(0, A4[1] - 0.6 * cm, A4[0], 0.6 * cm, fill=1, stroke=0)
        canvas.setFillColor(gold)
        canvas.rect(0, A4[1] - 0.7 * cm, A4[0], 0.1 * cm, fill=1, stroke=0)
        # Footer
        canvas.setFillColor(blue)
        canvas.setFont("Helvetica", 8)
        canvas.drawString(2 * cm, 1 * cm, "NPAMS User Manual")
        canvas.drawCentredString(A4[0] / 2, 1 * cm, TODAY)
        canvas.drawRightString(A4[0] - 2 * cm, 1 * cm, f"Page {doc.page}")
        canvas.restoreState()

    pdf = SimpleDocTemplate(
        out_path, pagesize=A4,
        leftMargin=2 * cm, rightMargin=2 * cm,
        topMargin=2 * cm, bottomMargin=2 * cm,
        title="NPAMS User Manual",
        author="LanFrame",
    )
    pdf.build(story, onFirstPage=_on_page, onLaterPages=_on_page)


if __name__ == "__main__":
    here = os.path.dirname(os.path.abspath(__file__))
    out_dir = os.path.abspath(os.path.join(here, "..", "docs"))
    os.makedirs(out_dir, exist_ok=True)
    docx_path = os.path.join(out_dir, "NPAMS_User_Manual.docx")
    pdf_path = os.path.join(out_dir, "NPAMS_User_Manual.pdf")
    build_docx(docx_path)
    build_pdf(pdf_path)
    print("Wrote:", docx_path)
    print("Wrote:", pdf_path)
